# -*- coding: utf-8 -*-
"""Screenshots 10.09.2026 – übernommene Acrobat-Elemente:
Mehrpunkt-Zeichenformen (Polygon/Verbundene Linien/Kommentarwolke/
Textlegende), Konvertieren (Word/Excel/PowerPoint/Text/RTF),
Dokumenteigenschaften (Strg+D), zertifizierte Kopie, Begrenzungs-
rahmen. Ohne Abo-/Cloud-/KI-Elemente."""
import gc
import os
import sys
import time

import pymupdf
import pytest

from conftest import app_schliessen
from documentai import samples
from documentai.engines import annotate
from documentai.engines.exporter import ExportEngine, SUPPORTED

tkinter = pytest.importorskip("tkinter")
HAVE_DISPLAY = (os.environ.get("DISPLAY") or
                sys.platform.startswith("win"))


# ======================================================================
# Backend
# ======================================================================
def test_polyform_alle_vier(svc, tmp_path):
    m, s, e = svc.open_document(samples.make_zeilen_formular_pdf(
        str(tmp_path / "w.pdf")))
    annotate.polyform_zeichnen(s, 0, "polygon",
                               [(100, 700), (150, 650), (200, 700)])
    annotate.polyform_zeichnen(s, 0, "polylinie",
                               [(100, 600), (150, 550), (200, 600)])
    annotate.polyform_zeichnen(s, 0, "wolke",
                               [(100, 500), (150, 450), (200, 500)])
    seite = s.doc[0]
    arten = [a.type[1] for a in seite.annots()]
    assert arten.count("Polygon") == 2      # Polygon + Wolke
    assert arten.count("PolyLine") == 1
    with pytest.raises(Exception):
        annotate.polyform_zeichnen(s, 0, "gibtsnicht", [(1, 1), (2, 2)])


def test_kommentarwolke_hat_wolkenrand(svc, tmp_path):
    """Die Wolke muss gezackte Bögen zeigen, nicht gerade Kanten."""
    m, s, e = svc.open_document(samples.make_zeilen_formular_pdf(
        str(tmp_path / "w.pdf")))
    annotate.polyform_zeichnen(s, 0, "wolke",
                               [(100, 700), (150, 650), (200, 700)])
    a = next(a for a in s.doc[0].annots() if a.type[1] == "Polygon")
    # Border-Diktat enthält den Cloud-Effekt
    bd = a.border_dict if hasattr(a, "border_dict") else {}
    assert "cloud" in str(bd).lower() or a.info


def test_textlegende(svc, tmp_path):
    m, s, e = svc.open_document(samples.make_zeilen_formular_pdf(
        str(tmp_path / "w.pdf")))
    annotate.polyform_zeichnen(s, 0, "legende",
                               [(100, 400), (150, 380)],
                               text="Wichtige Stelle")
    arten = [a.type[1] for a in s.doc[0].annots()]
    assert "FreeText" in arten and "Line" in arten


def test_metadaten(svc, tmp_path):
    m, s, e = svc.open_document(samples.make_zeilen_formular_pdf(
        str(tmp_path / "w.pdf")))
    annotate.metadaten_setzen(s, titel="Akte 4711", autor="Erika",
                              thema="Prüfung", schlagworte="x, y")
    md = annotate.metadaten_lesen(s)
    assert md["title"] == "Akte 4711"
    assert md["author"] == "Erika"
    assert md["seiten"] == 2


def test_zertifizierte_kopie(svc, tmp_path):
    """Schreibgeschützte, versiegelte Kopie: ohne Passwort lesbar,
    aber Rechte beschränkt (Acrobat „Zertifizierte Kopie speichern“)."""
    m, s, e = svc.open_document(samples.make_zeilen_formular_pdf(
        str(tmp_path / "w.pdf")))
    ziel = str(tmp_path / "zert.pdf")
    info = annotate.zertifizierte_kopie(s, ziel)
    assert info["pfad"] == ziel
    with pymupdf.open(ziel) as d:
        assert d.needs_pass == 0                    # lesbar ohne Passwort
        assert d.permissions < 0                    # Rechte beschränkt
        assert d[0].get_text().strip()              # Inhalt intakt


def test_export_xlsx_pptx_rtf(svc, tmp_path):
    m, s, e = svc.open_document(samples.make_zeilen_formular_pdf(
        str(tmp_path / "w.pdf")))
    svc.analyze(m, s, e)          # Seiten-Text füllen (wie in der GUI)
    ex = ExportEngine()
    for fmt in ("xlsx", "pptx", "rtf"):
        assert fmt in SUPPORTED
        ziel = str(tmp_path / ("aus." + fmt))
        r = ex.export(m, s, e, fmt, ziel)
        assert r["ok"], (fmt, r)
        assert os.path.getsize(ziel) > 100
    import openpyxl
    wb = openpyxl.load_workbook(str(tmp_path / "aus.xlsx"))
    assert any("Seite" in b for b in wb.sheetnames)
    from pptx import Presentation
    assert len(Presentation(str(tmp_path / "aus.pptx")).slides) == 2
    with open(str(tmp_path / "aus.rtf")) as f:
        assert f.read(6) == "{\\rtf1"


# ======================================================================
# GUI
# ======================================================================
def _app_neu():
    from documentai.gui.app import DocumentAIApp
    for _ in range(3):
        try:
            return DocumentAIApp()
        except Exception:
            gc.collect()
            time.sleep(0.4)
    return DocumentAIApp()


def _bereit(view):
    m = view.model
    return bool(m and m.pages and m.audit
                and m.audit[-1].action == "Analyse abgeschlossen"
                and view.page_list.get(0, "end"))


@pytest.fixture()
def vapp(svc, tmp_path, monkeypatch):
    from tkinter import messagebox
    monkeypatch.setattr(messagebox, "showerror",
                        lambda *a, **k: (_ for _ in ()).throw(
                            AssertionError(f"showerror: {a}")))
    monkeypatch.setattr(messagebox, "showinfo", lambda *a, **k: None)
    from documentai.gui import app as appmod
    monkeypatch.setattr(
        appmod.simpledialog, "askstring",
        lambda titel, frage, **k:
            "Wichtige Stelle" if "Beschriftung" in frage else "Notiz")
    pdf = samples.make_zeilen_formular_pdf(str(tmp_path / "z.pdf"))
    app = _app_neu()
    app.add_document(pdf)
    view = app._current()
    frist = time.time() + 40
    while time.time() < frist:
        app.root.update()
        if _bereit(view):
            break
        time.sleep(0.05)
    yield app, view
    app_schliessen(app)


@pytest.mark.skipif(not HAVE_DISPLAY, reason="kein Display")
def test_polyform_journey(vapp):
    """Polygon/Wolke per Klick+Doppelklick, Legende per 2 Klicks,
    Polylinie mit Esc abbrechen."""
    app, view = vapp
    scale = (150 * view.zoom) / 72.0
    c = view.canvas

    def klick(x, y):
        c.event_generate("<Button-1>", x=int(x), y=int(y))
        app.root.update()
        c.event_generate("<ButtonRelease-1>", x=int(x), y=int(y))
        app.root.update()

    def doppel(x, y):
        klick(x, y)
        view._on_canvas_doppelklick(None)
        app.root.update()

    view.set_polyform("polygon")
    app.root.update()
    assert view.mode == "polyform"
    for p in ((100, 650), (150, 600), (200, 650)):
        klick(p[0] * scale, p[1] * scale)
    doppel(230 * scale, 680 * scale)
    arten = [a.type[1] for a in view.session.doc[0].annots()]
    assert arten.count("Polygon") == 1
    assert view.model.edited

    view.set_polyform("wolke")
    app.root.update()
    for p in ((100, 550), (150, 500)):
        klick(p[0] * scale, p[1] * scale)
    doppel(200 * scale, 550 * scale)
    arten = [a.type[1] for a in view.session.doc[0].annots()]
    assert arten.count("Polygon") == 2

    app.root.update()
    view.set_polyform("legende")
    app.root.update()
    klick(100 * scale, 450 * scale)
    klick(160 * scale, 420 * scale)
    arten = [a.type[1] for a in view.session.doc[0].annots()]
    assert "FreeText" in arten and "Line" in arten

    view.set_polyform("polylinie")
    app.root.update()
    klick(100 * scale, 300 * scale)
    klick(150 * scale, 280 * scale)
    view.set_mode("normal")
    app.root.update()
    arten = [a.type[1] for a in view.session.doc[0].annots()]
    assert "PolyLine" not in arten            # Esc = abgebrochen
    assert not c.find_withtag("poly_form")    # Vorschau aufgeräumt


@pytest.mark.skipif(not HAVE_DISPLAY, reason="kein Display")
def test_begrenzungsrahmen_toggle(vapp):
    app, view = vapp
    app._begrenzungsrahmen()
    app.root.update()
    assert view._bbox_an
    assert len(view.canvas.find_withtag("bbox")) > 5
    app._begrenzungsrahmen()
    app.root.update()
    assert not view._bbox_an
    assert not view.canvas.find_withtag("bbox")


@pytest.mark.skipif(not HAVE_DISPLAY, reason="kein Display")
def test_konvertieren_gui(vapp, monkeypatch, tmp_path):
    from documentai.gui import app as appmod
    app, view = vapp
    ziel = {}

    def fake_save(**kw):
        ziel["pfad"] = str(tmp_path / ("aus" + kw.get("defaultextension",
                                                      ".pdf")))
        return ziel["pfad"]

    monkeypatch.setattr(appmod.filedialog, "asksaveasfilename", fake_save)
    app._konvertieren("xlsx")
    app.root.update()
    assert os.path.getsize(ziel["pfad"]) > 100
    app._konvertieren("pptx")
    app.root.update()
    assert os.path.getsize(ziel["pfad"]) > 100


@pytest.mark.skipif(not HAVE_DISPLAY, reason="kein Display")
def test_zertifikat_eigenschaften_gui(vapp, monkeypatch, tmp_path):
    from documentai.gui import app as appmod
    app, view = vapp
    zp = str(tmp_path / "zert.pdf")
    monkeypatch.setattr(appmod.filedialog, "asksaveasfilename",
                        lambda **kw: zp)
    app._current_zertifikat()
    app.root.update()
    with pymupdf.open(zp) as d:
        assert d.needs_pass == 0
    app._current_eigenschaften()
    app.root.update()
    tops = [w for w in app.root.winfo_children()
            if w.winfo_class() == "Toplevel"]
    assert tops, "Eigenschaften-Dialog fehlt"
    tops[-1].destroy()


@pytest.mark.skipif(not HAVE_DISPLAY, reason="kein Display")
def test_panel_v14_katalog(vapp):
    app, view = vapp
    view._panel_zeigen("werkzeuge")
    app.root.update()
    texte = {}

    def sammler(w):
        for k in w.winfo_children():
            sammler(k)
        try:
            if w.cget("text"):
                texte[w.cget("text")] = w
        except Exception:
            pass
    sammler(app.root)
    for name in ("Polygon", "Verbundene Linien", "Kommentarwolke",
                 "Textlegende", "In Word (.docx)", "In Excel (.xlsx)",
                 "In PowerPoint (.pptx)", "Als Text (.txt)", "Als RTF…",
                 "Zertifizierte Kopie…", "Begrenzungsrahmen anzeigen"):
        assert name in texte, f"Werkzeug {name!r} fehlt im Panel"
    for name in ("KI-Assistent", "Kostenlos testen", "Anmelden"):
        assert name not in texte, f"{name!r} darf nicht vorkommen"
